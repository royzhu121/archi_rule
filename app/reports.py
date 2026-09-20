"""生成适合打印的 A4 Word 审图报告。"""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.models import ReviewReportRequest


def _set_font(run, name: str = "微软雅黑", size: int = 10, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    _set_font(paragraph.add_run(text), size=16 if level == 1 else 12, bold=True)


def build_review_report(payload: ReviewReportRequest) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run("租户二次装修消防审图结果"), size=18, bold=True)

    boundary = document.add_paragraph()
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(
        boundary.add_run(
            "固定适用边界：已建成投运、耐火等级一级、多层民用建筑大型商业综合体；"
            "既有消防设施完整；不改变防火分区/防烟分区主边界"
        ),
        size=9,
    )

    _add_heading(document, "一、项目基本信息", 1)
    request = payload.request
    table = document.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    values = [
        ("站点", request.site, "楼层", request.floor),
        ("业态", f"{request.tenant_type_1} > {request.tenant_type_2}", "面积", f"{request.area:g}㎡"),
        ("吊顶", request.ceiling_type, "净空高度", f"{request.ceiling_height:g}m"),
    ]
    for row, row_values in zip(table.rows, values):
        for index, value in enumerate(row_values):
            row.cells[index].text = value
            if index % 2 == 0:
                _set_cell_shading(row.cells[index], "D9EAF7")
            for run in row.cells[index].paragraphs[0].runs:
                _set_font(run, bold=index % 2 == 0)

    result = payload.result
    _add_heading(document, "二、总体结论", 1)
    document.add_paragraph(result.overall_summary)
    if result.manual_reasons:
        for reason in result.manual_reasons:
            document.add_paragraph(reason, style="List Bullet")
    if result.llm_opinion:
        _add_heading(document, "AI 辅助意见", 2)
        document.add_paragraph(result.llm_opinion)

    _add_heading(document, "三、各专业审查结果", 1)
    if not result.modules:
        document.add_paragraph("本次触发人工复核边界，未生成各专业自动审查结论。")
    for module in result.modules.values():
        _add_heading(document, module.module_name, 2)
        document.add_paragraph(f"结论：{module.summary}")
        for detail in module.details:
            if detail:
                document.add_paragraph(detail, style="List Bullet")
        if module.calculations:
            document.add_paragraph(
                "计算/参数：" + "；".join(f"{key}={value}" for key, value in module.calculations.items())
            )
        document.add_paragraph("整改建议：")
        if module.suggestions:
            for suggestion in module.suggestions:
                document.add_paragraph(suggestion, style="List Bullet")
        else:
            document.add_paragraph("未触发新增整改项；仍须按原设计图纸和现场条件复核。")
        if module.references:
            document.add_paragraph("规范依据（仓库核验摘要，非规范原文）：")
            for reference in module.references:
                document.add_paragraph(
                    f"{reference.standard} 第{reference.article}条《{reference.title}》：{reference.excerpt}",
                    style="List Bullet",
                )

    _add_heading(document, "四、免责声明", 1)
    document.add_paragraph(
        "本报告仅用于上述固定建筑边界内的租户二次装修快速辅助审查，不构成正式消防审查意见，"
        "不得用于其他建筑、独立商店或新建工程。图纸、现场条件、原系统余量及规范适用性仍须由"
        "具备相应能力的消防专业人员复核；触发人工复核的事项不得依据本报告直接施工。"
    )

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
