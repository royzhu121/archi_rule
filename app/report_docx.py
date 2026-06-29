"""Word 报告生成。"""

from io import BytesIO
import re

from docx import Document
from docx.shared import Pt

from app.models import ReviewRequest, ReviewResult


def _clean_line(text: str) -> str:
    # 清理前缀符号，避免报告里出现重复 bullet 或异常符号。
    cleaned = re.sub(r"^[\s\u2022\-\*✦⚠🚨ℹ┌├└→]+", "", text or "")
    return cleaned.strip()


def _add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v


def build_review_report_docx(req: ReviewRequest, result: ReviewResult) -> bytes:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    doc.add_heading("租户二次消防审图报告", level=1)
    _add_key_value_table(
        doc,
        [
            ("站点", req.site or ""),
            ("租户名称", req.tenant_name or ""),
            ("楼层", req.floor or ""),
            ("一级业态", req.tenant_type_1 or ""),
            ("二级业态", req.tenant_type_2 or ""),
            ("房间面积(㎡)", f"{req.area}"),
            ("净空高度(m)", f"{req.ceiling_height}"),
            ("吊顶类型", req.ceiling_type or ""),
        ],
    )

    doc.add_heading("总体结论", level=2)
    doc.add_paragraph(_clean_line(result.overall_summary))
    doc.add_paragraph(
        "本工具基于租户入驻不改变原有防火分区、防烟分区主边界，且不涉及中庭/步行街/异形大空间的租户二次消防图纸审核。"
    )
    doc.add_paragraph(
        "生成结论仅供参考，不代替人工审核，所有审核结果须经消防专业工程师复核确认，并应遵循规范条文及企业要求。"
    )

    if result.highlights:
        doc.add_heading("风险与提醒", level=2)
        for line in result.highlights:
            clean = _clean_line(line)
            if clean:
                doc.add_paragraph(clean)

    doc.add_heading("各专业模块结论", level=2)
    for module_key, mod in result.modules.items():
        _ = module_key
        doc.add_heading(mod.module_name, level=3)
        doc.add_paragraph(f"状态：{mod.status}")
        doc.add_paragraph(f"摘要：{_clean_line(mod.summary)}")

        if mod.details:
            doc.add_paragraph("要点：")
            for line in mod.details:
                clean = _clean_line(line)
                if clean:
                    doc.add_paragraph(clean)

        if mod.impacts:
            doc.add_paragraph("影响项：")
            for line in mod.impacts:
                clean = _clean_line(line)
                if clean:
                    doc.add_paragraph(clean)

        if mod.suggestions:
            doc.add_paragraph("整改建议：")
            for line in mod.suggestions:
                clean = _clean_line(line)
                if clean:
                    doc.add_paragraph(clean)

        if mod.references:
            doc.add_paragraph("规范依据：")
            for ref in mod.references:
                doc.add_paragraph(f"{ref.standard} 第{ref.article}条 {ref.title}")

    if result.llm_opinion:
        doc.add_heading("AI 综合审查意见", level=2)
        doc.add_paragraph(_clean_line(result.llm_opinion))

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
